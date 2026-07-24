from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_service import RetornoExcelService
from app.services.retorno_service import RetornoService
from tests.test_retorno_service import _registro, _valor17


def _recebimento(seq, txid, chave, data, orig, pago, tarifa, nome):
    return _registro(
        {
            1: "5",
            2: txid,
            37: "60701190",
            45: "02",
            47: "00012345678901",
            89: chave,
            166: "2",
            167: "06",
            169: data,
            199: _valor17(orig),
            301: _valor17(pago),
            728: _valor17(tarifa),
            334: "01",
            336: "00098765432100",
            350: nome,
            632: "E" + seq.rjust(30, "0"),
            745: seq.rjust(6, "0"),
        }
    )


def _arquivo():
    header = _registro(
        {
            1: "0",
            2: "2",
            3: "RETORNO",
            10: "02",
            12: "PIX",
            27: "60701190",
            35: "02",
            37: "00012345678901",
            156: "20240319",
            254: "EMPRESA RECEBEDORA LTDA",
            732: "0000000001",
            742: "002",
            745: "000001",
        }
    )
    r1 = _recebimento("000002", "TX1", "a@x.com", "20240420", "1000.50", "1000.50", "1.90", "FULANO")
    r2 = _recebimento("000003", "TX2", "a@x.com", "20240420", "200.00", "200.00", "1.90", "BELTRANO")
    r3 = _recebimento("000004", "TX3", "b@y.com", "20240421", "300.25", "300.25", "1.90", "CICLANO")
    trailer = _registro(
        {1: "9", 2: "2", 3: "02", 5: "60701190", 713: _valor17("1500.75"), 730: "000000000000005", 745: "000005"}
    )
    conteudo = "\r\n".join([header, r1, r2, r3, trailer])
    return RetornoService.retorno_to_json(conteudo)


def _abrir():
    dados = RetornoExcelService.gerar_excel(_arquivo())
    return load_workbook(BytesIO(dados))


def test_gera_xlsx_com_abas():
    wb = _abrir()
    assert wb.sheetnames == ["Resumo", "Recebimentos", "Receita por Dia", "Receita por Chave"]


def test_detalhe_recebimentos():
    ws = _abrir()["Recebimentos"]
    assert ws.max_row - 1 == 3  # 3 recebimentos
    # valor pago (coluna L) e tarifa (coluna M) escritos como números
    pagos = [ws[f"L{r}"].value for r in range(2, 5)]
    assert pagos == [1000.50, 200.00, 300.25]
    assert ws["N2"].value == "=L2-M2"  # receita líquida


def test_resumo_usa_formulas():
    ws = _abrir()["Resumo"]
    formulas = {
        c.value
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    }
    assert "=SUM(Recebimentos!$L$2:$L$4)" in formulas  # receita bruta
    assert "=SUM(Recebimentos!$M$2:$M$4)" in formulas  # tarifas
    assert "=SUM(Recebimentos!$N$2:$N$4)" in formulas  # receita líquida
    assert any("IFERROR" in f for f in formulas)       # ticket médio


def test_receita_por_dia_agrupa_datas():
    ws = _abrir()["Receita por Dia"]
    datas = [ws[f"A{r}"].value for r in range(2, ws.max_row)]  # exceto TOTAL
    assert datas == ["2024-04-20", "2024-04-21"]
    assert ws["C2"].value == "=SUMIFS(Recebimentos!$L$2:$L$4,Recebimentos!$B$2:$B$4,$A2)"
    assert ws[f"A{ws.max_row}"].value == "TOTAL"


def test_receita_por_chave_agrupa_chaves():
    ws = _abrir()["Receita por Chave"]
    chaves = [ws[f"A{r}"].value for r in range(2, ws.max_row)]
    assert chaves == ["a@x.com", "b@y.com"]
    assert ws["C2"].value == "=SUMIFS(Recebimentos!$L$2:$L$4,Recebimentos!$C$2:$C$4,$A2)"
