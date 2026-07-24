"""Testes do processamento em streaming (parsing e geração de Excel).

Cobrem o splitter incremental (nos dois modos e com fronteiras de bloco
arbitrárias), a paridade entre o parsing em streaming e o não-streaming, a
equivalência da planilha gerada em streaming e os endpoints correspondentes.
"""

from io import BytesIO

from openpyxl import load_workbook

from app.schemas.retorno import RetornoRecebimento
from app.services.excel_service import RetornoExcelService
from app.services.retorno_service import RetornoService, _SplitterRegistros
from tests.test_excel_service import _recebimento
from tests.test_retorno_service import _registro, _valor17


def _conteudo(sep: str) -> str:
    header = _registro(
        {
            1: "0",
            2: "2",
            3: "RETORNO",
            10: "02",
            12: "PIX",
            27: "60701190",
            35: "02",
            37: "00012345678901",
            156: "20240319",
            254: "EMPRESA RECEBEDORA LTDA",
            732: "0000000001",
            742: "002",
            745: "000001",
        }
    )
    r1 = _recebimento(
        "000002", "TX1", "a@x.com", "20240420", "1000.50", "1000.50", "1.90", "FULANO"
    )
    r2 = _recebimento(
        "000003", "TX2", "a@x.com", "20240420", "200.00", "200.00", "1.90", "BELTRANO"
    )
    r3 = _recebimento(
        "000004", "TX3", "b@y.com", "20240421", "300.25", "300.25", "1.90", "CICLANO"
    )
    trailer = _registro(
        {
            1: "9",
            2: "2",
            3: "02",
            5: "60701190",
            713: _valor17("1500.75"),
            730: "000000000000005",
            745: "000005",
        }
    )
    return sep.join([header, r1, r2, r3, trailer])


# ---------------------------------------------------------------- #
# Splitter incremental                                             #
# ---------------------------------------------------------------- #
def _split_em_pedacos(dados: bytes, tamanho: int):
    sp = _SplitterRegistros()
    saida = []
    for i in range(0, len(dados), tamanho):
        saida.extend(sp.feed(dados[i : i + tamanho]))
    saida.extend(sp.flush())
    return saida


def test_splitter_modo_linha_varias_fronteiras():
    dados = _conteudo("\r\n").encode("utf-8")
    esperado = [r for r in _conteudo("\r\n").split("\r\n")]
    for tam in (1, 7, 250, 751, 4096, len(dados) + 10):
        regs = _split_em_pedacos(dados, tam)
        assert len(regs) == 5, f"tam={tam}"
        assert [r[0] for r in regs] == ["0", "5", "5", "5", "9"]
        assert all(len(r) == 750 for r in regs)
        assert regs == esperado


def test_splitter_modo_fixo_sem_quebra():
    dados = _conteudo("").encode("utf-8")  # blocos concatenados de 750
    for tam in (1, 13, 750, 1000, len(dados)):
        regs = _split_em_pedacos(dados, tam)
        assert len(regs) == 5, f"tam={tam}"
        assert [r[0] for r in regs] == ["0", "5", "5", "5", "9"]
        assert all(len(r) == 750 for r in regs)


def test_splitter_lida_com_lf_simples():
    dados = _conteudo("\n").encode("utf-8")
    regs = _split_em_pedacos(dados, 33)
    assert [r[0] for r in regs] == ["0", "5", "5", "5", "9"]


# ---------------------------------------------------------------- #
# Paridade streaming x não-streaming                               #
# ---------------------------------------------------------------- #
def test_stream_json_igual_ao_nao_stream():
    for sep in ("\r\n", "\n", ""):
        conteudo = _conteudo(sep)
        a = RetornoService.retorno_to_json(conteudo)
        b = RetornoService.retorno_to_json_stream(BytesIO(conteudo.encode("utf-8")))
        assert a.model_dump(mode="json") == b.model_dump(mode="json"), sep


def test_stream_json_recebimentos():
    b = RetornoService.retorno_to_json_stream(
        BytesIO(_conteudo("\r\n").encode("utf-8"))
    )
    recs = [d for d in b.detalhes if isinstance(d, RetornoRecebimento)]
    assert len(recs) == 3
    assert str(recs[0].valor_pago) == "1000.50"
    assert b.header.nome_recebedor == "EMPRESA RECEBEDORA LTDA"


def test_stream_json_vazio_levanta():
    import pytest

    with pytest.raises(ValueError):
        RetornoService.retorno_to_json_stream(BytesIO(b""))


# ---------------------------------------------------------------- #
# Resumo agregado em streaming                                     #
# ---------------------------------------------------------------- #
def test_resumo_stream_totais():
    r = RetornoService.resumo_stream(BytesIO(_conteudo("\r\n").encode("utf-8")))
    assert r.quantidade == 3
    assert str(r.receita_bruta) == "1500.75"
    assert str(r.tarifas) == "5.70"
    assert str(r.receita_liquida) == "1495.05"
    assert str(r.ticket_medio) == "500.25"
    assert r.nome_recebedor == "EMPRESA RECEBEDORA LTDA"
    assert r.ispb_participante == "60701190"


def test_resumo_stream_por_dia_e_chave():
    r = RetornoService.resumo_stream(BytesIO(_conteudo("").encode("utf-8")))
    dias = {g.chave: g for g in r.por_dia}
    assert set(dias) == {"2024-04-20", "2024-04-21"}
    assert dias["2024-04-20"].quantidade == 2
    assert str(dias["2024-04-20"].valor_pago) == "1200.50"
    assert str(dias["2024-04-21"].receita_liquida) == "298.35"

    chaves = {g.chave: g for g in r.por_chave}
    assert set(chaves) == {"a@x.com", "b@y.com"}
    assert chaves["a@x.com"].quantidade == 2
    assert str(chaves["b@y.com"].valor_pago) == "300.25"


def test_resumo_stream_independe_do_separador():
    a = RetornoService.resumo_stream(BytesIO(_conteudo("\r\n").encode("utf-8")))
    b = RetornoService.resumo_stream(BytesIO(_conteudo("").encode("utf-8")))
    assert a.model_dump(mode="json") == b.model_dump(mode="json")


# ---------------------------------------------------------------- #
# Excel agregado (a partir do resumo)                              #
# ---------------------------------------------------------------- #
def _excel_resumo():
    r = RetornoService.resumo_stream(BytesIO(_conteudo("\r\n").encode("utf-8")))
    return load_workbook(BytesIO(RetornoExcelService.gerar_excel_resumo(r)))


def test_excel_resumo_abas():
    wb = _excel_resumo()
    assert wb.sheetnames == ["Resumo", "Receita por Dia", "Receita por Chave"]


def test_excel_resumo_valores():
    wb = _excel_resumo()
    ws = wb["Resumo"]
    valores = {
        ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
        for r in range(1, ws.max_row + 1)
    }
    assert valores["Receita bruta (valor pago)"] == 1500.75
    assert valores["Tarifas de cobrança"] == 5.70
    assert valores["Receita líquida"] == 1495.05
    assert valores["Qtde. de recebimentos"] == 3


def test_excel_resumo_agrupamentos_e_total():
    wb = _excel_resumo()
    dia = wb["Receita por Dia"]
    datas = [dia[f"A{r}"].value for r in range(2, dia.max_row)]
    assert datas == ["2024-04-20", "2024-04-21"]
    assert dia[f"A{dia.max_row}"].value == "TOTAL"
    assert dia[f"C{dia.max_row}"].value == f"=SUM(C2:C{dia.max_row - 1})"

    chave = wb["Receita por Chave"]
    chaves = [chave[f"A{r}"].value for r in range(2, chave.max_row)]
    assert chaves == ["a@x.com", "b@y.com"]


# ---------------------------------------------------------------- #
# Endpoints (TestClient)                                           #
# ---------------------------------------------------------------- #
def _client():
    from fastapi.testclient import TestClient

    from app.main import app

    return TestClient(app)


def _upload(conteudo: bytes):
    return {"arquivo": ("retorno.rem", conteudo, "text/plain")}


def test_endpoint_retorno_to_json_stream():
    resp = _client().post(
        "/api/v1/retorno-to-json", files=_upload(_conteudo("\r\n").encode("utf-8"))
    )
    assert resp.status_code == 200
    data = resp.json()
    assert [d["tipo_registro"] for d in data["detalhes"]] == ["5", "5", "5"]
    assert data["header"]["nome_recebedor"] == "EMPRESA RECEBEDORA LTDA"


def test_endpoint_retorno_resumo():
    resp = _client().post(
        "/api/v1/retorno-resumo", files=_upload(_conteudo("").encode("utf-8"))
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["quantidade"] == 3
    assert data["receita_bruta"] == "1500.75"
    assert len(data["por_dia"]) == 2
    assert len(data["por_chave"]) == 2


def test_endpoint_excel_detalhado_para_arquivo_pequeno():
    resp = _client().post(
        "/api/v1/retorno-to-excel", files=_upload(_conteudo("\r\n").encode("utf-8"))
    )
    assert resp.status_code == 200
    assert resp.content[:2] == b"PK"
    wb = load_workbook(BytesIO(resp.content))
    # arquivo pequeno => planilha detalhada (inclui a aba Recebimentos)
    assert wb.sheetnames == [
        "Resumo",
        "Recebimentos",
        "Receita por Dia",
        "Receita por Chave",
    ]


def test_endpoint_excel_agregado_para_arquivo_grande(monkeypatch):
    import app.api.endpoints as ep

    # Força o caminho agregado sem precisar de um arquivo de 40 MB.
    monkeypatch.setattr(ep, "LIMITE_DETALHE_BYTES", 0)
    resp = _client().post(
        "/api/v1/retorno-to-excel", files=_upload(_conteudo("").encode("utf-8"))
    )
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    # sem a aba de detalhe
    assert wb.sheetnames == ["Resumo", "Receita por Dia", "Receita por Chave"]
