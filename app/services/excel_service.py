"""Geração de planilha Excel de análise de um arquivo de RETORNO CNAB750.

A análise foca na **receita** proveniente dos registros de recebimento
(detalhe tipo 5). São produzidas as abas:

- ``Resumo``            - totais consolidados (fórmulas sobre a aba de detalhe)
- ``Recebimentos``      - uma linha por recebimento
- ``Receita por Dia``   - sumarização por data de movimento
- ``Receita por Chave`` - sumarização por chave Pix do recebedor

Todos os totais são escritos como fórmulas (``SUM``/``SUMIFS``/``COUNTIFS``),
de modo que a planilha recalcula sozinha ao ser editada.
"""

from decimal import Decimal
from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..schemas.retorno import ArquivoRetorno, RetornoRecebimento

FONTE = "Arial"
MOEDA = '"R$" #,##0.00'

_TITULO_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
_TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
_BRANCO = Font(name=FONTE, bold=True, color="FFFFFF")
_BORDA_FINA = Border(bottom=Side(style="thin", color="BFBFBF"))

# Colunas da aba Recebimentos (1-based) usadas nas fórmulas das outras abas
_COL_DATA = "B"
_COL_CHAVE = "C"
_COL_VALOR_PAGO = "L"
_COL_TARIFA = "M"


def _num(valor: Optional[Decimal]) -> float:
    return float(valor) if valor is not None else 0.0


class RetornoExcelService:
    @staticmethod
    def gerar_excel(arquivo: ArquivoRetorno) -> bytes:
        recebimentos: List[RetornoRecebimento] = [
            d for d in arquivo.detalhes if isinstance(d, RetornoRecebimento)
        ]

        wb = Workbook()
        RetornoExcelService._aba_recebimentos(wb, recebimentos)
        RetornoExcelService._aba_resumo(wb, arquivo, recebimentos)
        RetornoExcelService._aba_por_dia(wb, recebimentos)
        RetornoExcelService._aba_por_chave(wb, recebimentos)

        # Deixa a aba Resumo em primeiro e como ativa
        ordem = ["Resumo", "Recebimentos", "Receita por Dia", "Receita por Chave"]
        wb._sheets.sort(key=lambda s: ordem.index(s.title))
        wb.active = 0

        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------ #
    # Aba: Recebimentos (detalhe)                                         #
    # ------------------------------------------------------------------ #
    @staticmethod
    def _aba_recebimentos(wb: Workbook, recebimentos: List[RetornoRecebimento]):
        ws = wb.active
        ws.title = "Recebimentos"

        colunas = [
            ("Identificador (TxId)", 22, None),
            ("Data Movimento", 15, None),
            ("Chave Pix", 26, None),
            ("Pagador Final", 30, None),
            ("CPF/CNPJ Pagador", 18, None),
            ("Valor Original", 15, MOEDA),
            ("Juros", 12, MOEDA),
            ("Multa", 12, MOEDA),
            ("Desconto", 12, MOEDA),
            ("Abatimento", 12, MOEDA),
            ("Valor Final", 15, MOEDA),
            ("Valor Pago", 15, MOEDA),
            ("Tarifa", 12, MOEDA),
            ("Receita Líquida", 15, MOEDA),
            ("End To End Id", 34, None),
            ("Cód. Liquidação", 14, None),
        ]
        for idx, (titulo, largura, _) in enumerate(colunas, start=1):
            letra = get_column_letter(idx)
            cel = ws.cell(row=1, column=idx, value=titulo)
            cel.font = _BRANCO
            cel.fill = _HEADER_FILL
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[letra].width = largura

        for i, r in enumerate(recebimentos, start=2):
            valores = [
                r.identificador,
                r.data_movimento.isoformat() if r.data_movimento else "",
                r.chave_pix,
                r.nome_pagador_final,
                r.cpf_cnpj_pagador_final,
                _num(r.valor_original),
                _num(r.valor_juros),
                _num(r.valor_multa),
                _num(r.valor_desconto),
                _num(r.valor_abatimento),
                _num(r.valor_final),
                _num(r.valor_pago),
                _num(r.tarifa_cobranca),
                None,  # receita líquida = valor pago - tarifa (fórmula abaixo)
                r.end_to_end_id,
                r.codigo_liquidacao,
            ]
            for idx, (valor, (_, _, fmt)) in enumerate(zip(valores, colunas), start=1):
                cel = ws.cell(row=i, column=idx, value=valor)
                cel.font = Font(name=FONTE)
                cel.border = _BORDA_FINA
                if fmt:
                    cel.number_format = fmt
            # Receita líquida (coluna N) = Valor Pago (L) - Tarifa (M)
            ws.cell(row=i, column=14, value=f"={_COL_VALOR_PAGO}{i}-{_COL_TARIFA}{i}")
            ws.cell(row=i, column=14).number_format = MOEDA

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(1, len(recebimentos) + 1)}"

    # ------------------------------------------------------------------ #
    # Aba: Resumo                                                         #
    # ------------------------------------------------------------------ #
    @staticmethod
    def _aba_resumo(
        wb: Workbook,
        arquivo: ArquivoRetorno,
        recebimentos: List[RetornoRecebimento],
    ):
        ws = wb.create_sheet("Resumo")
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 20

        # Faixa de dados na aba Recebimentos (linha 2 até a última)
        ultima = len(recebimentos) + 1 if recebimentos else 2

        def faixa(col: str) -> str:
            return f"Recebimentos!${col}$2:${col}${ultima}"

        # Título
        ws.merge_cells("A1:B1")
        titulo = ws["A1"]
        titulo.value = "Análise de Retorno CNAB750 — Receita"
        titulo.font = Font(name=FONTE, bold=True, size=14, color="FFFFFF")
        titulo.fill = _TITULO_FILL
        titulo.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 26

        # Metadados do arquivo
        info = [
            ("ISPB do recebedor", arquivo.header.ispb_participante or "-"),
            ("Recebedor", arquivo.header.nome_recebedor or "-"),
            (
                "Data de geração",
                arquivo.header.data_geracao.isoformat()
                if arquivo.header.data_geracao
                else "-",
            ),
        ]
        linha = 3
        for rotulo, valor in info:
            ws.cell(row=linha, column=1, value=rotulo).font = Font(name=FONTE, bold=True)
            ws.cell(row=linha, column=2, value=valor).font = Font(name=FONTE)
            linha += 1

        # Indicadores (fórmulas sobre a aba Recebimentos)
        linha += 1
        cab = ws.cell(row=linha, column=1, value="Indicador")
        cab.font = _BRANCO
        cab.fill = _HEADER_FILL
        cabv = ws.cell(row=linha, column=2, value="Valor")
        cabv.font = _BRANCO
        cabv.fill = _HEADER_FILL
        cabv.alignment = Alignment(horizontal="right")
        linha += 1

        indicadores = [
            ("Qtde. de recebimentos", f"=COUNTA({faixa('A')})", "0"),
            ("Valor original total", f"=SUM({faixa('F')})", MOEDA),
            ("Juros recebidos", f"=SUM({faixa('G')})", MOEDA),
            ("Multa recebida", f"=SUM({faixa('H')})", MOEDA),
            ("Descontos concedidos", f"=SUM({faixa('I')})", MOEDA),
            ("Abatimentos concedidos", f"=SUM({faixa('J')})", MOEDA),
            ("Receita bruta (valor pago)", f"=SUM({faixa('L')})", MOEDA),
            ("Tarifas de cobrança", f"=SUM({faixa('M')})", MOEDA),
            ("Receita líquida", f"=SUM({faixa('N')})", MOEDA),
            ("Ticket médio (valor pago)", None, MOEDA),
        ]
        primeira_ind = linha
        for rotulo, formula, fmt in indicadores:
            c_rot = ws.cell(row=linha, column=1, value=rotulo)
            c_rot.font = Font(name=FONTE)
            c_rot.border = _BORDA_FINA
            c_val = ws.cell(row=linha, column=2)
            c_val.font = Font(name=FONTE)
            c_val.border = _BORDA_FINA
            c_val.number_format = fmt
            c_val.alignment = Alignment(horizontal="right")
            if formula:
                c_val.value = formula
            linha += 1

        # Ticket médio = receita bruta / qtde (guardando divisão por zero)
        qtde_cel = f"B{primeira_ind}"
        bruta_cel = f"B{primeira_ind + 6}"
        ticket_cel = ws.cell(row=linha - 1, column=2)
        ticket_cel.value = f"=IFERROR({bruta_cel}/{qtde_cel},0)"

        # Destaca a linha de receita líquida
        rl = primeira_ind + 8
        for col in (1, 2):
            ws.cell(row=rl, column=col).font = Font(name=FONTE, bold=True)
            ws.cell(row=rl, column=col).fill = _TOTAL_FILL

    # ------------------------------------------------------------------ #
    # Aba: Receita por Dia                                                #
    # ------------------------------------------------------------------ #
    @staticmethod
    def _aba_por_dia(wb: Workbook, recebimentos: List[RetornoRecebimento]):
        chaves = sorted(
            {
                r.data_movimento.isoformat()
                for r in recebimentos
                if r.data_movimento is not None
            }
        )
        RetornoExcelService._aba_sumarizada(
            wb,
            titulo="Receita por Dia",
            rotulo_coluna="Data Movimento",
            largura_coluna=16,
            coluna_criterio=_COL_DATA,
            chaves=chaves,
            total_recebimentos=len(recebimentos),
        )

    # ------------------------------------------------------------------ #
    # Aba: Receita por Chave Pix                                          #
    # ------------------------------------------------------------------ #
    @staticmethod
    def _aba_por_chave(wb: Workbook, recebimentos: List[RetornoRecebimento]):
        chaves = sorted({r.chave_pix for r in recebimentos if r.chave_pix})
        RetornoExcelService._aba_sumarizada(
            wb,
            titulo="Receita por Chave",
            rotulo_coluna="Chave Pix",
            largura_coluna=30,
            coluna_criterio=_COL_CHAVE,
            chaves=chaves,
            total_recebimentos=len(recebimentos),
        )

    @staticmethod
    def _aba_sumarizada(
        wb: Workbook,
        titulo: str,
        rotulo_coluna: str,
        largura_coluna: int,
        coluna_criterio: str,
        chaves: List[str],
        total_recebimentos: int,
    ):
        ws = wb.create_sheet(titulo)
        ultima = total_recebimentos + 1 if total_recebimentos else 2

        cabecalhos = [
            (rotulo_coluna, largura_coluna, None),
            ("Qtde.", 10, "0"),
            ("Valor Pago", 16, MOEDA),
            ("Tarifas", 14, MOEDA),
            ("Receita Líquida", 16, MOEDA),
        ]
        for idx, (nome, largura, _) in enumerate(cabecalhos, start=1):
            cel = ws.cell(row=1, column=idx, value=nome)
            cel.font = _BRANCO
            cel.fill = _HEADER_FILL
            cel.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(idx)].width = largura

        crit = f"Recebimentos!${coluna_criterio}$2:${coluna_criterio}${ultima}"
        rng_pago = f"Recebimentos!${_COL_VALOR_PAGO}$2:${_COL_VALOR_PAGO}${ultima}"
        rng_tar = f"Recebimentos!${_COL_TARIFA}$2:${_COL_TARIFA}${ultima}"

        linha = 2
        for chave in chaves:
            ws.cell(row=linha, column=1, value=chave).font = Font(name=FONTE)
            ws.cell(row=linha, column=2, value=f"=COUNTIFS({crit},$A{linha})")
            ws.cell(row=linha, column=3, value=f"=SUMIFS({rng_pago},{crit},$A{linha})")
            ws.cell(row=linha, column=4, value=f"=SUMIFS({rng_tar},{crit},$A{linha})")
            ws.cell(row=linha, column=5, value=f"=C{linha}-D{linha}")
            for col in range(2, 6):
                ws.cell(row=linha, column=col).font = Font(name=FONTE)
                if col >= 3:
                    ws.cell(row=linha, column=col).number_format = MOEDA
            linha += 1

        # Linha de total
        if chaves:
            ws.cell(row=linha, column=1, value="TOTAL").font = Font(name=FONTE, bold=True)
            for col, letra in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
                cel = ws.cell(row=linha, column=col, value=f"=SUM({letra}2:{letra}{linha - 1})")
                cel.font = Font(name=FONTE, bold=True)
                cel.fill = _TOTAL_FILL
                if col >= 3:
                    cel.number_format = MOEDA
            ws.cell(row=linha, column=2).number_format = "0"

        ws.freeze_panes = "A2"
